import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from ..errors import PipelineError


def _is_group(shape) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.GROUP
    except NotImplementedError:
        return False


def _shape_texts(shape):
    """Recurse into groups and tables, which hasattr(shape, 'text') silently skipped."""
    if _is_group(shape):
        for s in shape.shapes:
            yield from _shape_texts(s)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text_frame.text.strip() for c in row.cells if c.text_frame.text.strip()]
            if cells:
                yield ", ".join(cells)
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            yield t


def extract_ppt_content(path: Path) -> list[dict]:
    try:
        prs = Presentation(str(path))
    except Exception:
        raise PipelineError("Could not read this presentation. It may be corrupt.")

    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        shapes = sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0))  # top-to-bottom reading order
        body = "\n".join(t for s in shapes for t in _shape_texts(s))

        slides.append({
            "index": idx,
            "hidden": slide._element.get("show") == "0",
            "notes": notes,
            "body": body,
        })
    return slides


def extract_pdf_content(path: Path) -> list[dict]:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise PipelineError("This PDF is password-protected.")
        return [
            {"index": i + 1, "hidden": False, "notes": "", "body": (p.extract_text() or "").strip()}
            for i, p in enumerate(reader.pages)
        ]
    except PipelineError:
        raise
    except Exception:
        raise PipelineError("Could not read this PDF. It may be corrupt.")


def clean_for_speech(text: str) -> str:
    lines = [re.sub(r"\s+", " ", ln).strip(" •-–*\t") for ln in text.splitlines()]
    lines = [ln for ln in lines if ln]
    # Add pauses between bullets so TTS doesn't run them together
    return " ".join(ln if ln[-1] in ".!?:;,।" else ln + "." for ln in lines)


def pick_script(slide: dict, mode: str = "auto") -> str:
    if mode == "notes":
        text = slide["notes"]
    elif mode == "slide":
        text = slide["body"]
    else:  # auto: speaker notes first, then slide text
        text = slide["notes"] or slide["body"]
    return clean_for_speech(text)